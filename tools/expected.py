# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl==3.1.5"]
# ///
"""Expected answers for the C# tests, computed independently of the C# code.

Reads the sample workbooks with openpyxl (a different reader), finds cells by scanning labels (not by
the server's layout constants), and computes the margin bridge with plain arithmetic. The C# tests
compare the server's results with this file.

    uv run tools/expected.py
"""

import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "samples" / "approved-exports"
OUT = ROOT / "tests" / "FinancialStatements.Tests" / "expected.json"
COMPANY = "Alder & Finch Supply Co."
UNITS = ["Wholesale", "Online", "Service"]


def cell(file, sheet, row_label, column_label):
    ws = load_workbook(DATA / file)[sheet]
    col = next(c.column for r in ws.iter_rows() for c in r if c.value == column_label)
    row = next(c.row for c in ws["A"] if c.value == row_label)
    c = ws.cell(row=row, column=col)
    return c.value, f"[{file}]{sheet}!{c.coordinate}"


def main():
    q = "Quarter to Date Actual"
    out = {"values": [], "bridge": {}}
    for file, sheet, row in [("IS_2026-06.xlsx", "Total", "Net sales"), ("IS_2026-06.xlsx", "Wholesale", "Gross profit"),
                             ("IS_2026-03.xlsx", "Total", "Gross margin %"), ("IS_2026-06.xlsx", "Service", "Freight billed to customers")]:
        v, ref = cell(file, sheet, row, q)
        out["values"].append({"file": file, "sheet": sheet, "row": row, "column": q, "value": v, "ref": ref})

    def figures(file):
        f = {}
        for u in UNITS:
            f[u] = (cell(file, u, "Gross profit", q)[0], cell(file, u, "Net sales", q)[0])
        f["Total"] = (cell(file, "Total", "Gross profit", q)[0], cell(file, "Total", "Net sales", q)[0])
        return f

    a, b = figures("IS_2026-03.xlsx"), figures("IS_2026-06.xlsx")
    sna, snb = sum(a[u][1] for u in UNITS), sum(b[u][1] for u in UNITS)
    Ma_units = sum(a[u][0] for u in UNITS) / sna
    Mb_units = sum(b[u][0] for u in UNITS) / snb
    units = {}
    for u in UNITS:
        ma, mb = a[u][0] / a[u][1], b[u][0] / b[u][1]
        wa, wb = a[u][1] / sna, b[u][1] / snb
        units[u] = {"rate_from": ma, "rate_to": mb, "rate_effect": wb * (mb - ma), "mix_effect": (wb - wa) * (ma - Ma_units)}
    out["bridge"] = {
        "company_rate_from": a["Total"][0] / a["Total"][1],
        "company_rate_to": b["Total"][0] / b["Total"][1],
        "units": units,
        "rate_effect_total": sum(x["rate_effect"] for x in units.values()),
        "mix_effect_total": sum(x["mix_effect"] for x in units.values()),
        "units_change": Mb_units - Ma_units,
    }
    OUT.write_text(json.dumps(out, indent=2) + "\n")
    print(json.dumps(out["bridge"], indent=1))


if __name__ == "__main__":
    main()
